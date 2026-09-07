"""
Export report snapshots to CSV and Excel.
"""

from __future__ import annotations

import csv
import io
from typing import Iterator

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reports.services.report_parser import ReportParser, flatten_rows
from reports.services.report_registry import ReportConfig


class ExportService:
    @staticmethod
    def _iter_rows(snapshot, config: ReportConfig) -> Iterator[list[str]]:
        columns = ReportParser.get_visible_columns(config)
        rows = flatten_rows(snapshot.raw_data)
        sum_row = snapshot.raw_data.get("SumRowData")

        yield [col["label"] for col in columns]

        for row in rows:
            yield [
                ReportParser.get_cell_display(row, col["key"], col["numeric"])
                for col in columns
            ]

        if sum_row:
            yield [
                ReportParser.get_sum_row_display(sum_row, col["key"], col["numeric"])
                if col["key"] in sum_row
                else ("جمع کل" if col == columns[0] else "—")
                for col in columns
            ]

    @staticmethod
    def to_csv(snapshot, config: ReportConfig) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        for row in ExportService._iter_rows(snapshot, config):
            writer.writerow(row)
        return buffer.getvalue().encode("utf-8-sig")

    @staticmethod
    def to_excel(snapshot, config: ReportConfig) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = config.title[:31]
        ws.sheet_view.rightToLeft = True

        columns = ReportParser.get_visible_columns(config)
        rows = flatten_rows(snapshot.raw_data)
        sum_row = snapshot.raw_data.get("SumRowData")

        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        sum_fill = PatternFill("solid", fgColor="E2EFDA")
        sum_font = Font(bold=True)

        headers = [col["label"] for col in columns]
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row in rows:
            ws.append(
                [
                    ReportParser.get_cell_display(row, col["key"], col["numeric"])
                    for col in columns
                ]
            )

        if sum_row:
            sum_values = []
            for i, col in enumerate(columns):
                if i == 0:
                    sum_values.append("جمع کل")
                elif col["key"] in sum_row:
                    sum_values.append(
                        ReportParser.get_sum_row_display(
                            sum_row, col["key"], col["numeric"]
                        )
                    )
                else:
                    sum_values.append("—")
            ws.append(sum_values)
            for cell in ws[ws.max_row]:
                cell.fill = sum_fill
                cell.font = sum_font

        for i, col in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(len(col["label"]) + 4, 14)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
